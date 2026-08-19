"""md → docx 转换(简历用):标题/粗体/列表/引用。"""
import re
import sys

from docx import Document
from docx.shared import Pt


def add_runs(paragraph, text: str, base_bold: bool = False) -> None:
    """解析 **bold** 和 [text](url) 生成 runs。"""
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*|\[(.+?)\]\((.+?)\)", text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()]).bold = base_bold
        if m.group(1) is not None:
            paragraph.add_run(m.group(1)).bold = True
        else:
            paragraph.add_run(m.group(2)).bold = base_bold
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:]).bold = base_bold


def convert(md_path: str, out_path: str) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    with open(md_path, encoding="utf-8") as f:
        for line in f:
            line = line.rstrip()
            if not line.strip():
                continue
            if line.strip() == "---":
                continue
            if line.startswith("#### "):
                doc.add_heading(line[5:].strip(), level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:].strip(), level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:].strip(), level=1)
            elif line.startswith("# "):
                doc.add_heading(line[2:].strip(), level=0)
            elif line.startswith("> "):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(18)
                add_runs(p, line[2:].strip())
                p.runs and [r.italic for r in p.runs]
            elif line.startswith("- "):
                p = doc.add_paragraph(style="List Bullet")
                add_runs(p, line[2:].strip())
            else:
                p = doc.add_paragraph()
                add_runs(p, line.strip())
    doc.save(out_path)
    print(f"OK -> {out_path}")


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
