"""Task 76.1 — POST /api/files session attachments + tenant isolation."""

from __future__ import annotations

import io

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from apps.api.routers.files import router as files_router
from packages.attachments.store import MemoryAttachmentStore, set_attachment_store
from packages.auth.dual_auth import verify_human_or_legacy_key
from packages.auth.models import TenantContext


@pytest.fixture
def files_api(tmp_path, monkeypatch: pytest.MonkeyPatch):
    import apps.api.routers.files as files_mod

    store = MemoryAttachmentStore()
    set_attachment_store(store)
    monkeypatch.setattr(files_mod, "UPLOAD_DIR", tmp_path)
    audits: list[dict] = []
    monkeypatch.setattr(
        "apps.api.routers.files.write_audit_sync",
        lambda rec: audits.append(rec) or True,
    )

    app = FastAPI()
    app.include_router(files_router)

    def _as(tid: str, uid: str, role: str = "user") -> TestClient:
        app.dependency_overrides[verify_human_or_legacy_key] = lambda: TenantContext(
            tid, uid, role, [], False
        )
        return TestClient(app)

    yield _as, store, audits, tmp_path
    set_attachment_store(None)
    app.dependency_overrides.clear()


def test_upload_txt_returns_id_and_blocks(files_api) -> None:
    as_client, store, _, tmp_path = files_api
    client = as_client("t1", "u1")
    r = client.post(
        "/api/files",
        files={"file": ("note.txt", b"hello contract\n", "text/plain")},
        data={"session_id": "s1"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "ready"
    assert body["blocks"]
    assert "hello" in body["blocks"][0]["text"]
    aid = body["attachment_id"]
    got = client.get(f"/api/files/{aid}/blocks", params={"session_id": "s1"})
    assert got.status_code == 200
    assert got.json()["blocks"]


def test_cross_tenant_blocks_403_and_audit(files_api) -> None:
    as_client, _, audits, _ = files_api
    owner = as_client("t1", "u1")
    aid = owner.post(
        "/api/files",
        files={"file": ("note.txt", b"secret\n", "text/plain")},
        data={"session_id": "s1"},
    ).json()["attachment_id"]
    other = as_client("t2", "u1")
    r = other.get(f"/api/files/{aid}/blocks", params={"session_id": "s1"})
    assert r.status_code == 403
    assert r.json()["detail"]["code"] == "AUTH_004"
    assert any(a.get("action") == "attachment_denied" for a in audits)


def test_xlsx_upload_table_blocks(files_api) -> None:
    from openpyxl import Workbook

    as_client, _, _, _ = files_api
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "name"
    ws["B1"] = "amt"
    ws["A2"] = "a"
    ws["B2"] = 1
    buf = io.BytesIO()
    wb.save(buf)
    client = as_client("t1", "u1")
    r = client.post(
        "/api/files",
        files={
            "file": (
                "t.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"session_id": "s1"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["blocks"][0]["kind"] == "table"


def test_docx_paragraphs(files_api) -> None:
    import docx as docx_mod

    buf = io.BytesIO()
    d = docx_mod.Document()
    d.add_paragraph("违约条款")
    d.save(buf)
    as_client, _, _, _ = files_api
    r = as_client("t1", "u1").post(
        "/api/files",
        files={
            "file": (
                "c.docx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"session_id": "s1"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "ready"
    assert "违约" in r.json()["blocks"][0]["text"]
    from pypdf import PdfWriter

    buf = io.BytesIO()
    w = PdfWriter()
    w.add_blank_page(width=72, height=72)
    w.write(buf)
    as_client, _, _, _ = files_api
    c = as_client("t1", "u1")
    r = c.post(
        "/api/files",
        files={"file": ("scan.pdf", buf.getvalue(), "application/pdf")},
        data={"session_id": "s1"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "ocr_required"
    assert r.json()["blocks"] == []
