"""Task 52 S3 — API helpers, Excel, router smoke."""

from __future__ import annotations

from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock, patch

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from packages.auth.dual_auth import verify_human_or_legacy_key
from packages.auth.models import TenantContext
from packages.social.exceptions import SocialCostAlertError
from packages.social.export_xlsx import HEADERS, build_analysis_xlsx
from packages.social.service import (
    RETRY_ANALYZE_ONLY,
    create_analysis_task,
    hint_previous_task_id,
    list_contents,
    resolve_brand,
    retry_task,
    stub_replica,
)
from packages.database.pgvector_session import (
    Base,
    SocialAccount,
    SocialContent,
    SocialFollow,
    SocialResult,
    SocialTask,
    SocialTemplate,
    SocialUsage,
    TenantConfig,
)
from apps.api.routers.social import router as social_router


@pytest.fixture()
def db_session():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    tables = [
        SocialAccount.__table__,
        SocialFollow.__table__,
        SocialContent.__table__,
        SocialTask.__table__,
        SocialTemplate.__table__,
        SocialResult.__table__,
        SocialUsage.__table__,
        TenantConfig.__table__,
    ]
    Base.metadata.create_all(engine, tables=tables)
    Session = sessionmaker(bind=engine)
    with Session() as session:
        yield session


def test_build_analysis_xlsx_meta_and_headers():
    raw = build_analysis_xlsx(
        rows=[
            {
                "platform": "douyin",
                "title": "t",
                "content": "文案" * 30,
                "external_id": "vid1",
                "duration_s": 12,
                "like_count": 3,
                "comment_count": None,
                "share_count": 1,
                "collect_count": 2,
                "published_at": datetime(2026, 1, 1, tzinfo=UTC),
                "fetched_at": datetime(2026, 1, 2, tzinfo=UTC),
                "content_source": "desc",
                "replica_json": {"titles": ["a"], "script": "s", "tags": ["x"]},
            }
        ],
        new_count=7,
    )
    wb = load_workbook(BytesIO(raw))
    ws = wb.active
    assert ws.cell(1, 1).value == "本次新增 7"
    assert [c.value for c in ws[2]] == HEADERS
    assert ws.freeze_panes == "A3"
    assert ws.column_dimensions["D"].width == 60
    assert ws.cell(3, 5).value == "vid1"
    assert ws.cell(3, 8).value in (None, "")  # missing comment_count stays empty


def test_resolve_brand_from_org_profile(db_session):
    db_session.add(
        TenantConfig(
            tenant_id="t1",
            config={
                "org_content_profile": {
                    "name": "星火",
                    "product_focus": "教育",
                    "target_audience": "家长",
                }
            },
        )
    )
    db_session.commit()
    brand = resolve_brand(db_session, "t1", None)
    assert brand["name"] == "星火"
    assert brand["business"] == "教育"
    assert brand["audience"] == "家长"
    assert "style" not in brand

    with pytest.raises(ValueError, match="企业画像"):
        resolve_brand(db_session, "empty", None)


def test_create_task_requires_prior_probe(db_session, monkeypatch):
    monkeypatch.setenv("SOCIAL_COST_ALERT_USD", "999")
    with pytest.raises(ValueError, match="请先拉取账号"):
        create_analysis_task(
            db_session,
            tenant_id="t1",
            user_id="u1",
            platform="douyin",
            account_key="never_probed",
        )


def test_create_task_duplicate_returns_existing(db_session, monkeypatch):
    monkeypatch.setenv("SOCIAL_COST_ALERT_USD", "999")
    acct = SocialAccount(platform="douyin", account_key="ak", external_id="SEC")
    db_session.add(acct)
    db_session.flush()
    t1, created = create_analysis_task(
        db_session,
        tenant_id="t1",
        user_id="u1",
        platform="douyin",
        account_key="ak",
    )
    assert created is True
    t2, created2 = create_analysis_task(
        db_session,
        tenant_id="t1",
        user_id="u1",
        platform="douyin",
        account_key="ak",
    )
    assert created2 is False
    assert t2.id == t1.id


def test_create_analyze_requires_contents(db_session, monkeypatch):
    monkeypatch.setenv("SOCIAL_COST_ALERT_USD", "999")
    acct = SocialAccount(platform="douyin", account_key="ak", external_id="SEC")
    db_session.add(acct)
    db_session.flush()
    with pytest.raises(ValueError, match="请先拉取内容"):
        create_analysis_task(
            db_session,
            tenant_id="t1",
            user_id="u1",
            platform="douyin",
            account_key="ak",
            phase="analyze",
        )


def test_create_fetch_phase_sets_marker(db_session, monkeypatch):
    monkeypatch.setenv("SOCIAL_COST_ALERT_USD", "999")
    acct = SocialAccount(platform="douyin", account_key="ak", external_id="SEC")
    db_session.add(acct)
    db_session.flush()
    task, created = create_analysis_task(
        db_session,
        tenant_id="t1",
        user_id="u1",
        platform="douyin",
        account_key="ak",
        phase="fetch",
    )
    assert created is True
    assert task.error == "[fetch_only:20]"


def test_create_fetch_phase_custom_limit(db_session, monkeypatch):
    monkeypatch.setenv("SOCIAL_COST_ALERT_USD", "999")
    acct = SocialAccount(platform="douyin", account_key="ak", external_id="SEC")
    db_session.add(acct)
    db_session.flush()
    task, created = create_analysis_task(
        db_session,
        tenant_id="t1",
        user_id="u1",
        platform="douyin",
        account_key="ak",
        phase="fetch",
        fetch_limit=35,
    )
    assert created is True
    assert task.error == "[fetch_only:35]"


def test_cost_alert_blocks_new_task(db_session, monkeypatch):
    monkeypatch.setenv("SOCIAL_COST_ALERT_USD", "0.01")
    acct = SocialAccount(platform="douyin", account_key="ak", external_id="SEC")
    db_session.add(acct)
    db_session.add(
        SocialUsage(
            tenant_id="t",
            user_id="u",
            platform="douyin",
            operation="probe",
            item_count=1,
            cost_usd=Decimal("1.0000"),
            price_usd=Decimal("3.0000"),
        )
    )
    db_session.flush()
    with pytest.raises(SocialCostAlertError):
        create_analysis_task(
            db_session,
            tenant_id="t1",
            user_id="u1",
            platform="douyin",
            account_key="ak",
        )


def test_hint_and_retry_analyze_only(db_session):
    acct = SocialAccount(platform="douyin", account_key="ak", external_id="SEC")
    db_session.add(acct)
    db_session.flush()
    prev = SocialTask(
        tenant_id="t1",
        user_id="u1",
        platform="douyin",
        account_id=acct.id,
        status="done",
        new_count=5,
    )
    cur = SocialTask(
        tenant_id="t1",
        user_id="u1",
        platform="douyin",
        account_id=acct.id,
        status="done",
        new_count=0,
    )
    db_session.add_all([prev, cur])
    db_session.flush()
    assert hint_previous_task_id(db_session, cur) == prev.id

    failed = SocialTask(
        tenant_id="t1",
        user_id="u1",
        platform="douyin",
        account_id=acct.id,
        status="failed",
        error="boom",
    )
    db_session.add(failed)
    db_session.add(
        SocialContent(
            platform="douyin",
            account_id=acct.id,
            external_id="v1",
            content="x",
        )
    )
    db_session.flush()
    retry_task(db_session, failed)
    assert failed.status == "pending"
    assert failed.error == RETRY_ANALYZE_ONLY


def test_stub_replica_requires_template(db_session, monkeypatch):
    monkeypatch.setenv("LLM_PROVIDER", "mock")
    acct = SocialAccount(platform="douyin", account_key="ak", external_id="SEC")
    db_session.add(acct)
    db_session.flush()
    task = SocialTask(
        tenant_id="t1",
        user_id="u1",
        platform="douyin",
        account_id=acct.id,
        status="done",
    )
    content = SocialContent(
        platform="douyin",
        account_id=acct.id,
        external_id="v1",
        content="body",
    )
    db_session.add_all([task, content])
    db_session.flush()
    with pytest.raises(ValueError, match="暂无可用模版"):
        stub_replica(
            db_session,
            tenant_id="t1",
            user_id="u1",
            content_id=content.id,
            task_id=task.id,
            brand={"name": "星火", "business": "教育"},
        )


def test_stub_replica_with_structure_precipitates_template(db_session, monkeypatch):
    monkeypatch.setenv("LLM_PROVIDER", "mock")
    acct = SocialAccount(platform="douyin", account_key="ak", external_id="SEC")
    db_session.add(acct)
    db_session.flush()
    task = SocialTask(
        tenant_id="t1",
        user_id="u1",
        platform="douyin",
        account_id=acct.id,
        status="done",
    )
    content = SocialContent(
        platform="douyin",
        account_id=acct.id,
        external_id="v1",
        content="口播" * 40,
        title="t",
    )
    result = SocialResult(
        task_id=0,  # set after flush
        content_id=0,
        structure_json={
            "template_type": "清单式",
            "hooks": ["痛点提问"],
            "outline": [
                {"role": "hook", "text": "提问"},
                {"role": "method", "text": "三点"},
                {"role": "cta", "text": "关注"},
            ],
            "topics": ["学习"],
            "replicables": ["三点法"],
        },
    )
    db_session.add_all([task, content])
    db_session.flush()
    result.task_id = task.id
    result.content_id = content.id
    db_session.add(result)
    db_session.flush()
    res = stub_replica(
        db_session,
        tenant_id="t1",
        user_id="u1",
        content_id=content.id,
        task_id=task.id,
        brand={"name": "星火", "business": "教育"},
    )
    assert res.template_id is not None
    assert res.replica_json is not None
    assert len(res.replica_json.get("titles") or []) == 3


def test_list_contents_filters(db_session):
    acct = SocialAccount(platform="douyin", account_key="ak", external_id="SEC")
    db_session.add(acct)
    db_session.flush()
    db_session.add_all(
        [
            SocialContent(
                platform="douyin",
                account_id=acct.id,
                external_id="a",
                like_count=10,
                collect_count=1,
                duration_s=30,
                content="yes",
            ),
            SocialContent(
                platform="douyin",
                account_id=acct.id,
                external_id="b",
                like_count=1,
                collect_count=0,
                duration_s=5,
                content="",
            ),
        ]
    )
    db_session.flush()
    rows = list_contents(db_session, account_id=acct.id, like_min=5, has_content=True)
    assert len(rows) == 1
    assert rows[0].external_id == "a"
    assert list_contents(db_session, account_id=99999) == []


def test_process_analyze_only_skips_tikhub(db_session, monkeypatch):
    from packages.social import pipeline as pl

    monkeypatch.setenv("LLM_PROVIDER", "mock")

    acct = SocialAccount(platform="douyin", account_key="ak", external_id="SEC")
    db_session.add(acct)
    db_session.flush()
    content = SocialContent(
        platform="douyin",
        account_id=acct.id,
        external_id="v1",
        content="口播文案",
    )
    task = SocialTask(
        tenant_id="t1",
        user_id="u1",
        platform="douyin",
        account_id=acct.id,
        status="running",
        error=RETRY_ANALYZE_ONLY,
    )
    db_session.add_all([content, task])
    db_session.commit()

    class BoomConn:
        def fetch_recent(self, *a, **k):
            raise AssertionError("must not call TikHub on analyze-only retry")

        def close(self) -> None:
            pass

    pl.process_task(
        db_session,
        db_session.query(SocialTask).one(),
        connector=BoomConn(),  # type: ignore[arg-type]
        acquire=lambda **_: True,
    )
    task = db_session.query(SocialTask).one()
    assert task.status == "done"
    assert db_session.query(SocialResult).count() == 1


def _tenant() -> TenantContext:
    return TenantContext("t1", "u1", "user", ["chat:write"], False)


def test_router_get_task_tenant_isolation(db_session):
    acct = SocialAccount(platform="douyin", account_key="ak", external_id="SEC")
    db_session.add(acct)
    db_session.flush()
    task = SocialTask(
        tenant_id="other",
        user_id="u9",
        platform="douyin",
        account_id=acct.id,
        status="done",
    )
    db_session.add(task)
    db_session.commit()

    class _SessCtx:
        def __enter__(self):
            return db_session

        def __exit__(self, *a):
            return False

    app = FastAPI()
    app.include_router(social_router)

    async def _auth() -> TenantContext:
        return _tenant()

    app.dependency_overrides[verify_human_or_legacy_key] = _auth
    with patch("apps.api.routers.social._db_session", lambda: _SessCtx()):
        with patch("apps.api.routers.social.log_audit", MagicMock()):
            client = TestClient(app)
            r = client.get(f"/api/social/analysis-tasks/{task.id}")
            assert r.status_code == 404


def test_probe_adds_tenant_follow(db_session, monkeypatch):
    from packages.social.service import list_followed_accounts, probe_account
    from packages.social.types import AccountInfo

    monkeypatch.setenv("SOCIAL_COST_ALERT_USD", "999")

    class FakeConn:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def probe(self, platform: str, account_key: str) -> AccountInfo:
            return AccountInfo(
                platform="douyin",
                account_key=account_key,
                external_id="SEC",
                nickname="老丁",
            )

    monkeypatch.setattr("packages.social.service.TikHubConnector", lambda: FakeConn())
    monkeypatch.setattr("packages.social.service.acquire_for_probe", lambda: True)

    row = probe_account(
        db_session,
        tenant_id="t1",
        user_id="u1",
        platform="douyin",
        account_key="ak",
    )
    mine = list_followed_accounts(db_session, tenant_id="t1")
    other = list_followed_accounts(db_session, tenant_id="t2")
    assert [a.id for a in mine] == [row.id]
    assert other == []


def test_router_follows_tenant_isolation(db_session):
    acct = SocialAccount(platform="douyin", account_key="ak", external_id="SEC")
    db_session.add(acct)
    db_session.flush()
    db_session.add(SocialFollow(tenant_id="other", account_id=acct.id))
    db_session.commit()

    class _SessCtx:
        def __enter__(self):
            return db_session

        def __exit__(self, *a):
            return False

    app = FastAPI()
    app.include_router(social_router)

    async def _auth() -> TenantContext:
        return _tenant()

    app.dependency_overrides[verify_human_or_legacy_key] = _auth
    with patch("apps.api.routers.social._db_session", lambda: _SessCtx()):
        with patch("apps.api.routers.social.log_audit", MagicMock()):
            client = TestClient(app)
            r = client.get("/api/social/follows")
            assert r.status_code == 200
            assert r.json() == {"items": []}
