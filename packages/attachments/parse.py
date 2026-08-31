"""Local document parsers → attachment_blocks (no embeddings)."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field

from packages.attachments.errors import ParseError
from packages.attachments.validate import validate_attachment_bytes

PAGE_CHARS = 800
TABLE_ROWS = 40


@dataclass
class AttachmentBlock:
    block_index: int
    kind: str
    text: str
    char_count: int
    page: int | None = None
    sheet: str | None = None
    rows: int | None = None


@dataclass
class ParseResult:
    status: str
    media_type: str
    blocks: list[AttachmentBlock] = field(default_factory=list)


def _chunk_text(text: str, *, kind: str, page: int | None = None) -> list[dict]:
    s = (text or "").strip()
    if not s:
        return []
    overlap = max(1, int(PAGE_CHARS * 0.1))
    out: list[dict] = []
    i = 0
    while i < len(s):
        piece = s[i : i + PAGE_CHARS]
        out.append({"kind": kind, "text": piece, "page": page})
        if i + PAGE_CHARS >= len(s):
            break
        i += PAGE_CHARS - overlap
    return out


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "gbk", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _parse_txt(data: bytes) -> list[dict]:
    return _chunk_text(_decode_text(data), kind="text")


def _parse_csv(data: bytes) -> list[dict]:
    text = _decode_text(data)
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return []
    header = rows[0]
    body = rows[1:]
    chunks: list[dict] = []
    for start in range(0, max(len(body), 1), TABLE_ROWS):
        group = body[start : start + TABLE_ROWS] or [[]]
        lines = [",".join(header)]
        lines.extend(",".join(r) for r in group)
        chunks.append(
            {
                "kind": "table",
                "text": "\n".join(lines),
                "sheet": "csv",
                "rows": len(group),
            }
        )
    return chunks


def _parse_pdf(data: bytes) -> tuple[str, list[dict]]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    chunks: list[dict] = []
    any_text = False
    for i, page in enumerate(reader.pages, start=1):
        raw = (page.extract_text() or "").strip()
        if raw:
            any_text = True
            chunks.extend(_chunk_text(raw, kind="page", page=i))
    if not any_text:
        return "ocr_required", []
    return "ready", chunks


def _parse_docx(data: bytes) -> list[dict]:
    import docx  # type: ignore[import-untyped]

    doc = docx.Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        lines = []
        for row in table.rows:
            lines.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
        if lines:
            parts.append("\n".join(lines))
    return _chunk_text("\n".join(parts), kind="text")


def _row_group_chunk(
    header: list[str], rows: list[list[str]], *, sheet: str
) -> dict:
    lines = [",".join(header)]
    lines.extend(",".join(r) for r in rows)
    return {
        "kind": "table",
        "text": "\n".join(lines),
        "sheet": sheet,
        "rows": len(rows),
    }


def _parse_xlsx(data: bytes) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    chunks: list[dict] = []
    try:
        for sheet in wb.worksheets:
            rows_iter = sheet.iter_rows(values_only=True)
            try:
                header = next(rows_iter)
            except StopIteration:
                continue
            header_s = ["" if c is None else str(c) for c in header]
            buf: list[list[str]] = []
            title = sheet.title
            for row in rows_iter:
                buf.append(["" if c is None else str(c) for c in row])
                if len(buf) >= TABLE_ROWS:
                    chunks.append(_row_group_chunk(header_s, buf, sheet=title))
                    buf = []
            if buf:
                chunks.append(_row_group_chunk(header_s, buf, sheet=title))
    finally:
        wb.close()
    return chunks


def _image_media_type(data: bytes) -> str:
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if data.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if data.startswith(b"GIF87a") or data.startswith(b"GIF89a"):
        return "image/gif"
    if len(data) >= 12 and data.startswith(b"RIFF") and data[8:12] == b"WEBP":
        return "image/webp"
    return "image/jpeg"


def parse_document(data: bytes, filename: str) -> ParseResult:
    kind = validate_attachment_bytes(filename=filename, data=data)
    media = {
        "pdf": "application/pdf",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "csv": "text/csv",
        "txt": "text/plain",
        "image": _image_media_type(data),
    }.get(kind, "application/octet-stream")
    try:
        if kind == "image":
            return ParseResult(status="ready", media_type=media, blocks=[])
        if kind == "pdf":
            status, raw = _parse_pdf(data)
        elif kind == "docx":
            status, raw = "ready", _parse_docx(data)
        elif kind == "xlsx":
            status, raw = "ready", _parse_xlsx(data)
        elif kind == "csv":
            status, raw = "ready", _parse_csv(data)
        else:
            status, raw = "ready", _parse_txt(data)
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("FILE_PARSE", str(exc)[:200]) from exc

    if kind != "pdf" and not raw:
        status = "failed"
    blocks = [
        AttachmentBlock(
            block_index=i,
            kind=str(item.get("kind") or "text"),
            text=str(item.get("text") or ""),
            char_count=len(str(item.get("text") or "")),
            page=item.get("page"),  # type: ignore[arg-type]
            sheet=item.get("sheet"),  # type: ignore[arg-type]
            rows=item.get("rows"),  # type: ignore[arg-type]
        )
        for i, item in enumerate(raw)
    ]
    if status == "ocr_required":
        return ParseResult(status=status, media_type=media, blocks=[])
    if status == "failed":
        return ParseResult(status=status, media_type=media, blocks=[])
    return ParseResult(status="ready", media_type=media, blocks=blocks)
