"""Excel export for social analysis (Task 52 S3)."""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

HEADERS = [
    "序号",
    "平台",
    "标题",
    "视频文案",
    "视频ID",
    "时长(s)",
    "点赞数",
    "评论数",
    "分享数",
    "收藏数",
    "发布时间",
    "采集时间",
    "发布时长",
    "视频链接",
    "文稿来源",
    "复刻稿",
]


def _fmt_dt(v: datetime | None) -> str:
    if v is None:
        return ""
    if v.tzinfo is None:
        return v.isoformat(sep=" ", timespec="seconds")
    return v.astimezone(UTC).strftime("%Y-%m-%d %H:%M:%S")


def _publish_age(published_at: datetime | None, fetched_at: datetime | None) -> str:
    if published_at is None:
        return ""
    end = fetched_at or datetime.now(UTC)
    if published_at.tzinfo is None and end.tzinfo is not None:
        published_at = published_at.replace(tzinfo=UTC)
    if end.tzinfo is None and published_at.tzinfo is not None:
        end = end.replace(tzinfo=UTC)
    delta = end - published_at
    days = delta.days
    hours = delta.seconds // 3600
    return f"{days}天{hours}小时"


def _video_url(platform: str, external_id: str) -> str:
    if platform == "douyin" and external_id:
        return f"https://www.douyin.com/video/{external_id}"
    return ""


def _replica_text(replica_json: dict[str, Any] | None) -> str:
    if not replica_json:
        return ""
    titles = replica_json.get("titles") or []
    script = replica_json.get("script") or ""
    tags = replica_json.get("tags") or []
    parts = []
    if titles:
        parts.append("标题: " + " | ".join(str(t) for t in titles))
    if script:
        parts.append(str(script))
    if tags:
        parts.append("标签: " + " ".join(f"#{t}" for t in tags))
    return "\n".join(parts)


def build_analysis_xlsx(
    *,
    rows: list[dict[str, Any]],
    new_count: int,
) -> bytes:
    """Single sheet: title row「本次新增 N」+ header freeze/autofilter + wrap 文案."""
    wb = Workbook()
    ws = wb.active
    ws.title = "对标分析"

    # Row 1: meta banner (spec: 表头「本次新增 N」)
    ws.append([f"本次新增 {new_count}"] + [""] * (len(HEADERS) - 1))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws.cell(row=1, column=1).font = Font(bold=True, size=12)

    # Row 2: column headers
    ws.append(list(HEADERS))
    for cell in ws[2]:
        cell.font = Font(bold=True)
    ws.auto_filter.ref = f"A2:{get_column_letter(len(HEADERS))}{2 + max(len(rows), 0)}"
    ws.freeze_panes = "A3"

    for i, r in enumerate(rows, start=1):
        content = r.get("content") or ""
        replica = _replica_text(r.get("replica_json"))
        line = [
            i,
            r.get("platform") or "",
            r.get("title") or "",
            content,
            r.get("external_id") or "",
            r.get("duration_s") if r.get("duration_s") is not None else "",
            r.get("like_count") if r.get("like_count") is not None else "",
            r.get("comment_count") if r.get("comment_count") is not None else "",
            r.get("share_count") if r.get("share_count") is not None else "",
            r.get("collect_count") if r.get("collect_count") is not None else "",
            _fmt_dt(r.get("published_at")),
            _fmt_dt(r.get("fetched_at")),
            _publish_age(r.get("published_at"), r.get("fetched_at")),
            _video_url(str(r.get("platform") or ""), str(r.get("external_id") or "")),
            r.get("content_source") or "",
            replica,
        ]
        ws.append(line)
        row_idx = i + 2
        for col in (4, 16):
            cell = ws.cell(row=row_idx, column=col)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        approx = max(len(content) // 40, len(replica) // 40, 1)
        ws.row_dimensions[row_idx].height = min(200, 15 * min(approx, 12))

    for col_idx, _ in enumerate(HEADERS, start=1):
        letter = get_column_letter(col_idx)
        if col_idx == 4:
            ws.column_dimensions[letter].width = 60
        elif col_idx == 16:
            ws.column_dimensions[letter].width = 40
        elif col_idx in (3, 14):
            ws.column_dimensions[letter].width = 24
        else:
            ws.column_dimensions[letter].width = 12

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
